from docx import Document
from docx.shared import Pt, RGBColor, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d=Document('ref_default.docx')
NAVY=RGBColor(0x1F,0x38,0x64); TEAL=RGBColor(0x12,0x7A,0x8A); GREY=RGBColor(0x55,0x60,0x6E)
def font(st,name='Calibri',size=None,bold=None,color=None,italic=None):
    st.font.name=name
    rpr=st.element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),name)
    for a in ('w:asciiTheme','w:hAnsiTheme','w:cstheme','w:eastAsiaTheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    if size: st.font.size=Pt(size)
    if bold is not None: st.font.bold=bold
    if italic is not None: st.font.italic=italic
    if color is not None: st.font.color.rgb=color
S0=d.styles
class _S:
    def __getitem__(self,k):
        for st in S0:
            if st.name==k: return st
        raise KeyError(k)
    def add_style(self,*a): return S0.add_style(*a)
S=_S()
for nm in ('Normal','Body Text','First Paragraph','Compact'):
    try: st=S[nm]
    except KeyError: continue
    font(st,'Calibri',11)
    pf=st.paragraph_format; pf.space_after=Pt(6); pf.space_before=Pt(0); pf.line_spacing=1.15
    if nm in ('Body Text','First Paragraph'): pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
S['Compact'].paragraph_format.space_after=Pt(2); S['Compact'].font.size=Pt(10)
for lvl,size,sb,sa in ((1,15,18,6),(2,12.5,12,4),(3,11.5,10,3)):
    st=S[f'Heading {lvl}']; font(st,'Cambria',size,True,NAVY if lvl<3 else TEAL,False)
    st.paragraph_format.space_before=Pt(sb); st.paragraph_format.space_after=Pt(sa); st.paragraph_format.keep_with_next=True
    if lvl==1:
        pPr=st.element.get_or_add_pPr(); bdr=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom')
        for k,v in (('w:val','single'),('w:sz','6'),('w:space','2'),('w:color','1F3864')): b.set(qn(k),v)
        bdr.append(b); pPr.append(bdr)
for nm,kw in (('Image Caption',dict(size=9.5,italic=True,color=GREY)),('Table Caption',dict(size=9.5,italic=False,bold=True,color=NAVY)),('Caption',dict(size=9.5,italic=True,color=GREY))):
    try: st=S[nm]; font(st,'Calibri',**kw); st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER; st.paragraph_format.space_after=Pt(10)
    except KeyError: pass
try: S['Table Caption'].paragraph_format.keep_with_next=True; S['Table Caption'].paragraph_format.space_before=Pt(8)
except KeyError: pass
try:
    st=S['TOC Heading']; font(st,'Cambria',15,True,NAVY)
except KeyError:
    st=S.add_style('TOC Heading',WD_STYLE_TYPE.PARAGRAPH); font(st,'Cambria',15,True,NAVY)
st.element.set(qn('w:styleId'),'TOCHeading')
def pstyle(name,size,bold=False,color=NAVY,align=WD_ALIGN_PARAGRAPH.CENTER,sb=0,sa=6,italic=False,fname='Cambria'):
    try: st=S[name]
    except KeyError: st=S.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st.base_style=S['Normal']; font(st,fname,size,bold,color,italic)
    pf=st.paragraph_format; pf.alignment=align; pf.space_before=Pt(sb); pf.space_after=Pt(sa)
    return st
pstyle('Cover Institution',15,True,NAVY,sb=40,sa=2)
pstyle('Cover Course',11.5,False,GREY,sa=70,fname='Calibri')
pstyle('Cover Title',24,True,NAVY,sa=14)
pstyle('Cover Subtitle',13,False,TEAL,sa=18,fname='Calibri')
pstyle('Cover Target',12,False,NAVY,sa=70,italic=True,fname='Calibri')
pstyle('Cover Meta',11.5,False,RGBColor(0x22,0x22,0x22),sa=6,fname='Calibri')
pstyle('Cover Note',9,False,GREY,sb=60,sa=0,italic=True,fname='Calibri')
r=pstyle('Reference',10,False,RGBColor(0x22,0x22,0x22),align=WD_ALIGN_PARAGRAPH.LEFT,sa=4,fname='Calibri')
r.paragraph_format.left_indent=Cm(0.9); r.paragraph_format.first_line_indent=Cm(-0.9)
p=pstyle('Prompt',10.5,False,RGBColor(0x1d,0x2b,0x3a),align=WD_ALIGN_PARAGRAPH.JUSTIFY,sa=6,italic=False,fname='Calibri')
p.paragraph_format.left_indent=Cm(0.8); p.paragraph_format.right_indent=Cm(0.8)
pPr=p.element.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),'EEF4F8'); pPr.append(shd)
bdr=OxmlElement('w:pBdr'); l=OxmlElement('w:left')
for k,v in (('w:val','single'),('w:sz','18'),('w:space','8'),('w:color','127A8A')): l.set(qn(k),v)
bdr.append(l); pPr.append(bdr)
# page setup + header/footer
sec=d.sections[0]
sec.page_width=Cm(21.0); sec.page_height=Cm(29.7)
sec.left_margin=sec.right_margin=Cm(2.3); sec.top_margin=Cm(2.3); sec.bottom_margin=Cm(2.0)
sec.different_first_page_header_footer=True
hp=sec.header.paragraphs[0]; hp.text='AI-Assisted Reverse Engineering of an Axial-Flow Compressor — Project Proposal (Phase 1)'
hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for run in hp.runs: run.font.size=Pt(8.5); run.font.color.rgb=GREY; run.font.name='Calibri'
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
def fld(par,code):
    r1=par.add_run(); f1=OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'),'begin'); r1._r.append(f1)
    r2=par.add_run(); it=OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text=code; r2._r.append(it)
    r3=par.add_run(); f3=OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'),'separate'); r3._r.append(f3)
    r4=par.add_run('1')
    r5=par.add_run(); f5=OxmlElement('w:fldChar'); f5.set(qn('w:fldCharType'),'end'); r5._r.append(f5)
    for r_ in (r1,r2,r3,r4,r5): r_.font.size=Pt(9); r_.font.color.rgb=GREY
a=fp.add_run('Turbomachinery · Amirkabir University of Technology · Bayati & Bali   |   Page '); a.font.size=Pt(9); a.font.color.rgb=GREY
fld(fp,' PAGE ')
b=fp.add_run(' of '); b.font.size=Pt(9); b.font.color.rgb=GREY
fld(fp,' NUMPAGES ')
d.save('reference.docx')
print('ref ok')
